"""Bounded local document extraction. No uploaded HTML/XML is executed."""
import csv
import io
import json
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from bs4 import BeautifulSoup
from defusedxml import ElementTree
from docx import Document
import openpyxl
import pypdf
import xlrd

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 200000
SUPPORTED = {'.pdf', '.docx', '.xlsx', '.xls', '.csv', '.json', '.xml', '.html', '.htm', '.md', '.markdown', '.txt'}


def decode_text(contents):
    if contents.startswith((b'\xff\xfe', b'\xfe\xff')):
        return contents.decode('utf-16')
    try:
        return contents.decode('utf-8-sig')
    except UnicodeDecodeError:
        return contents.decode('cp1251')


def extract_document(filename, contents):
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED:
        raise ValueError('Поддерживаются PDF, DOCX, XLSX, XLS, CSV, JSON, XML, HTML, Markdown и TXT. Старый DOC сохраните как DOCX.')
    if not contents or len(contents) > MAX_FILE_BYTES:
        raise ValueError('Файл пуст или больше 10 МБ.')
    if extension in {'.docx', '.xlsx'}:
        try:
            with ZipFile(io.BytesIO(contents)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 50*1024*1024:
                    raise ValueError('Слишком большой распакованный документ.')
        except BadZipFile as error:
            raise ValueError('Повреждённый документ Office.') from error
    parts, warnings = [], []
    character_count = 0

    def append(text):
        nonlocal character_count
        parts.append(str(text))
        character_count += len(parts[-1]) + 2
        if character_count > MAX_TEXT_CHARS:
            raise ValueError('В документе больше 200 000 символов. Разделите файл на части.')

    try:
        if extension == '.pdf':
            reader = pypdf.PdfReader(io.BytesIO(contents))
            if reader.is_encrypted:
                raise ValueError('PDF защищён паролем. Загрузите незашифрованную копию.')
            if len(reader.pages) > 300:
                raise ValueError('В PDF больше 300 страниц.')
            blank_pages = 0
            for page in reader.pages:
                text = page.extract_text() or ''
                blank_pages += int(not text.strip())
                append(text)
            if blank_pages:
                warnings.append(f'Страниц PDF без извлекаемого текста: {blank_pages}. OCR сканов не выполняется.')
        elif extension == '.docx':
            doc = Document(io.BytesIO(contents))
            for element in doc.element.body:
                if element.tag.endswith('}p'):
                    append(''.join(t.text or '' for t in element.iter() if t.tag.endswith('}t')))
                elif element.tag.endswith('}tbl'):
                    for row in element:
                        if row.tag.endswith('}tr'):
                            append(' | '.join(''.join(t.text or '' for t in cell.iter() if t.tag.endswith('}t')) for cell in row if cell.tag.endswith('}tc')))
        elif extension == '.xlsx':
            workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                        if row_index > 20000:
                            raise ValueError('В таблице слишком много строк.')
                        if any(value is not None for value in row):
                            append(' | '.join(str(value) for value in row if value is not None))
            finally:
                workbook.close()
            warnings.append('Excel: извлечены значения ячеек; формулы не пересчитываются.')
        elif extension == '.xls':
            workbook = xlrd.open_workbook(file_contents=contents)
            for sheet in workbook.sheets():
                if sheet.nrows > 20000:
                    raise ValueError('В таблице слишком много строк.')
                for index in range(sheet.nrows):
                    append(' | '.join(str(value) for value in sheet.row_values(index) if value != ''))
            workbook.release_resources()
        else:
            text = decode_text(contents)
            if len(text) > MAX_TEXT_CHARS:
                raise ValueError('В файле больше 200 000 символов.')
            if extension == '.json':
                def visit(value):
                    if isinstance(value, str):
                        append(value)
                    elif isinstance(value, list):
                        for item in value:
                            visit(item)
                    elif isinstance(value, dict):
                        for item in value.values():
                            visit(item)
                visit(json.loads(text))
            elif extension == '.xml':
                root = ElementTree.fromstring(text)
                for fragment in root.itertext():
                    if fragment.strip():
                        append(fragment.strip())
            elif extension in {'.html', '.htm'}:
                soup = BeautifulSoup(text, 'html.parser')
                for tag in soup(['script', 'style', 'noscript']):
                    tag.decompose()
                append(soup.get_text(' ', strip=True))
            elif extension == '.csv':
                try:
                    dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
                except csv.Error:
                    dialect = csv.excel
                for row in csv.reader(io.StringIO(text), dialect):
                    append(' | '.join(row))
            else:
                append(text)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f'Не удалось прочитать {extension}: файл повреждён или имеет неподдерживаемую структуру.') from error
    result = '\n\n'.join(p.strip() for p in parts if p.strip())
    if not result:
        raise ValueError('Не найден текст. Для PDF-сканов требуется предварительное OCR.')
    if len(result) > MAX_TEXT_CHARS:
        raise ValueError('Извлечено больше 200 000 символов. Разделите документ.')
    return {'text': result, 'filename': Path(filename).name, 'warnings': warnings}
