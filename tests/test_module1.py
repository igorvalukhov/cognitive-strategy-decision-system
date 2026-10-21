import io
import json
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from src.synthetic.dataset import parse_annotation, save_splits, is_near_duplicate
from src.causal_extraction.metrics import evaluate_sequences
from src.web.documents import extract_document
from src.web.app import create_app


def test_clean_offsets_and_complete_events():
    r = parse_annotation('Because <CAUSE>major donors lost confidence</CAUSE>, <EFFECT>the party cut its field staff</EFFECT>.', True)
    assert r['spans'][1]['text'] == 'the party cut its field staff'
    for s in r['spans']:
        assert r['text'][s['start']:s['end']] == s['text']
    assert r['labels'][1] == 'B-CAUSE'
    assert r['labels'][-2] == 'E-EFFECT'
    assert is_near_duplicate(r, [r])


@pytest.mark.parametrize('marked', [
    '<CAUSE>Donors left <EFFECT>the party</EFFECT></CAUSE> today.',
    '<CAUSE>Donors left</CAUSE> the party today.',
    '<CAUSE> Donors left </CAUSE> and <EFFECT>staff resigned</EFFECT>.'])
def test_bad_markup_rejected(marked):
    with pytest.raises(ValueError):
        parse_annotation(marked, True)


def test_stratified_splits_are_disjoint(tmp_path):
    records = [{'id':str(i), 'tokens':['word'], 'labels':['O'], 'causal':i >= 30} for i in range(100)]
    assert save_splits(records, tmp_path) == {'train':80, 'val':10, 'test':10}
    sets = [{json.loads(line)['id'] for line in (tmp_path/f'{s}.jsonl').read_text().splitlines()} for s in ('train','val','test')]
    assert not (sets[0]&sets[1] or sets[0]&sets[2] or sets[1]&sets[2])


def test_false_positives_on_o_reduce_precision():
    good = evaluate_sequences([[4,0,8,0]], [[4,0,8,0]])
    bad = evaluate_sequences([[4,0,8,4]], [[4,0,8,0]])
    assert good['exact_span']['f1'] == 1
    assert bad['exact_span']['precision'] == pytest.approx(2/3)
    assert bad['token_macro']['precision'] < good['token_macro']['precision']
    assert bad['exact_single_pair']['f1'] == 0
    negative = evaluate_sequences([[4,0,8]], [[0,0,0]])
    assert negative['negative_false_alarm_rate'] == 1


@pytest.mark.parametrize(('name','content'), [
    ('test.txt',b'Higher taxes caused protests.'),
    ('test.md',b'# Higher taxes caused protests.'),
    ('test.csv',b'text\nHigher taxes caused protests.'),
    ('test.json',b'{"text":"Higher taxes caused protests."}'),
    ('test.xml',b'<root>Higher taxes caused protests.</root>'),
    ('test.html',b'<p>Higher taxes caused protests.</p><script>unsafe()</script>')])
def test_text_formats(name,content):
    result = extract_document(name,content)
    assert 'Higher taxes caused protests.' in result['text']
    assert 'unsafe()' not in result['text']


def test_office_formats():
    from docx import Document
    import openpyxl
    doc = Document()
    doc.add_paragraph('Higher taxes caused protests.')
    stream = io.BytesIO()
    doc.save(stream)
    assert extract_document('a.docx',stream.getvalue())['text'] == 'Higher taxes caused protests.'
    book = openpyxl.Workbook()
    book.active.append(['Higher taxes caused protests.'])
    stream = io.BytesIO()
    book.save(stream)
    assert extract_document('a.xlsx',stream.getvalue())['text'] == 'Higher taxes caused protests.'


def test_reject_unsupported_and_xml_entities():
    with pytest.raises(ValueError):
        extract_document('a.doc',b'old binary')
    with pytest.raises(ValueError):
        extract_document('a.xml',b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///private">]><x>&x;</x>')


def test_upload_api():
    class Predictor:
        def predict_detailed(self,text):
            return {'causes':[], 'effects':[]}
    with TestClient(create_app(Predictor)) as client:
        result = client.post('/api/analyze-file',files={'file':('a.txt',b'Political parties met today.','text/plain')})
        assert result.status_code == 200
        assert result.json()['source'] == 'a.txt'
        assert result.json()['statistics']['total_sentences'] == 1
        assert client.post('/api/analyze-file',files={'file':('a.exe',b'no')}).status_code == 422


def test_first_subtoken_training_alignment(tmp_path):
    from src.causal_extraction.training_data import CausalDataset
    class Encoding(dict):
        def word_ids(self): return [None,0,0,1,None]
    class Tokenizer:
        pad_token_id = 0
        def __call__(self,*args,**kwargs):
            return Encoding(input_ids=[101,5,6,7,102],attention_mask=[1]*5)
    file = tmp_path/'train.jsonl'
    file.write_text(json.dumps({'tokens':['Unrestful','fall'],'labels':['S-CAUSE','S-EFFECT']}))
    data = CausalDataset(file,Tokenizer(),8)
    assert data[0]['labels'].tolist() == [-100,4,-100,8,-100,-100,-100,-100]
    with pytest.raises(ValueError): CausalDataset(file,Tokenizer(),4)


def test_generation_pipeline_with_test_double(tmp_path):
    from threading import Event
    from src.synthetic.generator import generate
    class Client:
        model, base_url = 'test-double', 'http://test.invalid'
        def health(self): return {'available':True}
        def complete(self,prompt,seed):
            if prompt.startswith('Check'): return {'valid':True}
            if 'must NOT assert' in prompt:
                return {'marked_text':f'Delegate{seed} read report{seed} about committee{seed} beside minister{seed} during session{seed}.'}
            return {'marked_text':f'<CAUSE>Policy{seed} reduced tax{seed}</CAUSE> so <EFFECT>company{seed} hired worker{seed}</EFFECT> in district{seed}.'}
    class Tokenizer:
        def __call__(self,tokens,**kwargs): return {'input_ids':[0]*(len(tokens)+2)}
    result = generate('test fixture',30,.3,42,tmp_path,Client(),Tokenizer(),lambda **kw:None,Event())
    assert result['count'] == 30
    assert sum(result['splits'].values()) == 30
    records = [json.loads(x) for x in (tmp_path/'accepted.jsonl').read_text().splitlines()]
    assert sum(not r['causal'] for r in records) == 9
    assert (tmp_path/'manifest.json').exists()
